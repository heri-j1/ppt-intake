#!/usr/bin/env python3
"""inspect_pptx.py — PPT → 结构 JSON 清单（ppt-intake skill Step 1）。

用法:
    python inspect_pptx.py <file.pptx> [-o out.json]

输出 JSON 包含: 画布尺寸 / 主题色与字体 / 每页角色推断 / 逐 shape
位置尺寸 / repr 级准确文本 / 字号 / 占位符类型 / 图片 / 表格 / 文本预算(budget)。
后续 merge plan 的 old_text 一律从本清单的 text 字段复制。
"""
import argparse
import json
import sys

from pptx import Presentation
from pptx.util import Emu

EMU_PER_CM = 360000.0


def cm(v):
    return round(Emu(v).cm, 2) if v is not None else None


def get_ea_font(rPr):
    """run 级 rPr 里的 <a:ea> 东亚字体（中文实际渲染用）。"""
    if rPr is None:
        return None
    ea = rPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}ea")
    return ea.get("typeface") if ea is not None else None


def text_shape_meta(shape):
    """文本形状 → 清单项。"""
    tf = shape.text_frame
    paragraphs = [p.text for p in tf.paragraphs]
    text = "\n".join(paragraphs)
    font_pt = None
    font_name = None
    font_ea = None
    for p in tf.paragraphs:
        for r in p.runs:
            if r.font.size is not None:
                pt = r.font.size.pt
                font_pt = max(font_pt or 0, pt)
            if r.font.name:
                font_name = r.font.name
            ea = get_ea_font(r._r.find(
                "{http://schemas.openxmlformats.org/drawingml/2006/main}rPr"))
            if ea:
                font_ea = ea
    orig_len = len(text.replace("\n", ""))
    h_cm = cm(shape.height) or 0
    # label/body 判定与 budget 规则见 rules/baseline-rules.md R4
    role_class = "label" if (h_cm < 1.5 or orig_len <= 8 or (font_pt or 0) >= 20) else "body"
    return {
        "id": shape.shape_id,
        "name": shape.name,
        "type": "TEXT",
        "placeholder": str(shape.placeholder_format.type) if shape.is_placeholder else None,
        "x_cm": cm(shape.left), "y_cm": cm(shape.top),
        "w_cm": cm(shape.width), "h_cm": cm(shape.height),
        "text": text,
        "paragraphs": paragraphs,
        "font_pt": font_pt, "font_name": font_name, "font_ea": font_ea,
        "orig_len": orig_len,
        "budget": int(orig_len * 1.1) + 1,
        "role_class": role_class,
    }


def shape_meta(shape):
    """任意形状 → 清单项（GROUP 递归）。"""
    try:
        if shape.shape_type == 6:  # GROUP
            return {
                "id": shape.shape_id, "name": shape.name, "type": "GROUP",
                "x_cm": cm(shape.left), "y_cm": cm(shape.top),
                "w_cm": cm(shape.width), "h_cm": cm(shape.height),
                "children": [shape_meta(c) for c in shape.shapes],
            }
        if getattr(shape, "has_table", False) and shape.has_table:
            return {
                "id": shape.shape_id, "name": shape.name, "type": "TABLE",
                "x_cm": cm(shape.left), "y_cm": cm(shape.top),
                "w_cm": cm(shape.width), "h_cm": cm(shape.height),
                "rows": [[c.text for c in row.cells] for row in shape.table.rows],
            }
        if shape.shape_type == 13:  # PICTURE
            meta = {
                "id": shape.shape_id, "name": shape.name, "type": "PICTURE",
                "x_cm": cm(shape.left), "y_cm": cm(shape.top),
                "w_cm": cm(shape.width), "h_cm": cm(shape.height),
            }
            try:
                img = shape.image
                meta["ext"] = img.ext
                meta["px"] = [img.size[0], img.size[1]]
            except Exception:
                meta["ext"] = None  # 外链或异常图片，克隆时走 rels 层不受影响
            return meta
        if getattr(shape, "has_chart", False) and shape.has_chart:
            return {
                "id": shape.shape_id, "name": shape.name, "type": "CHART",
                "x_cm": cm(shape.left), "y_cm": cm(shape.top),
                "w_cm": cm(shape.width), "h_cm": cm(shape.height),
                "chart_type": str(shape.chart.chart_type),
            }
        if shape.has_text_frame and shape.text_frame.text.strip():
            return text_shape_meta(shape)
        return {
            "id": shape.shape_id, "name": shape.name, "type": "SHAPE",
            "x_cm": cm(shape.left), "y_cm": cm(shape.top),
            "w_cm": cm(shape.width), "h_cm": cm(shape.height),
        }
    except Exception as e:  # 单形状解析失败不拖垮整份清单
        return {"id": getattr(shape, "shape_id", None), "name": getattr(shape, "name", "?"),
                "type": "ERROR", "error": str(e)}


def flatten(meta):
    if meta.get("type") == "GROUP":
        for c in meta.get("children", []):
            yield from flatten(c)
    else:
        yield meta


def infer_role(slide_idx, slide_count, shapes):
    """角色启发式（cover/section/content/stats/quote/closing），agent 应复核。"""
    texts = [m for m in flatten_shapes(shapes) if m.get("type") == "TEXT"]
    if not texts:
        return "blank"
    max_pt = max((t.get("font_pt") or 0) for t in texts)
    total_len = sum(t["orig_len"] for t in texts)
    big_num = any(t["font_pt"] and t["font_pt"] >= 54 and
                  any(ch.isdigit() for ch in t["text"]) for t in texts)
    tall_body = any((t.get("h_cm") or 0) >= 8 and t["orig_len"] > 20 for t in texts)
    if slide_idx == 0 and max_pt >= 36:
        return "cover"
    if big_num and len(texts) >= 3:
        return "stats"
    if slide_idx == slide_count - 1 and len(texts) <= 2 and total_len < 40:
        return "closing"
    if len(texts) <= 2 and total_len < 40 and not tall_body:
        return "section"
    if len(texts) == 1 and total_len > 60:
        return "quote"
    return "content"


def flatten_shapes(shapes):
    for m in shapes:
        yield from flatten(m)


def theme_info(prs):
    colors, fonts = {}, {}
    try:
        ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
        for part in prs.part.package.iter_parts():
            if "theme" not in str(part.partname):
                continue
            root = part._element if hasattr(part, "_element") else None
            if root is None:
                from lxml import etree
                root = etree.fromstring(part.blob)
            scheme = root.find(".//a:clrScheme", ns)
            if scheme is not None:
                for c in scheme:
                    tag = c.tag.split("}")[1]
                    srgb = c.find("a:srgbClr", ns)
                    sysc = c.find("a:sysClr", ns)
                    val = (srgb.get("val") if srgb is not None
                           else sysc.get("lastClr") if sysc is not None else None)
                    if val:
                        colors[tag] = "#" + val
            fs = root.find(".//a:fontScheme", ns)
            if fs is not None:
                for which in ("major", "minor"):
                    el = fs.find(f"a:{which}Font", ns)
                    if el is None:
                        continue
                    latin = el.find("a:latin", ns)
                    ea = el.find("a:ea", ns)
                    fonts[which] = {"latin": latin.get("typeface") if latin is not None else None,
                                    "ea": ea.get("typeface") if ea is not None else None}
            break  # 只取第一个 theme
    except Exception as e:
        colors["_error"] = str(e)
    return {"colors": colors, "fonts": fonts}


def inspect(path):
    prs = Presentation(path)
    slides = []
    for idx, slide in enumerate(prs.slides):
        shapes = [shape_meta(s) for s in slide.shapes]
        slides.append({
            "index": idx,
            "role": infer_role(idx, len(prs.slides), shapes),
            "layout": slide.slide_layout.name,
            "shapes": shapes,
        })
    return {
        "file": path,
        "slide_size_cm": [cm(prs.slide_width), cm(prs.slide_height)],
        "slide_count": len(prs.slides),
        "theme": theme_info(prs),
        "slides": slides,
    }


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("pptx")
    ap.add_argument("-o", "--out")
    args = ap.parse_args()
    data = inspect(args.pptx)
    text = json.dumps(data, ensure_ascii=False, indent=1)
    if args.out:
        with open(args.out, "w", encoding="utf-8") as f:
            f.write(text)
        print(f"OK {args.pptx}: {data['slide_count']} slides -> {args.out}")
    else:
        print(text)


if __name__ == "__main__":
    sys.exit(main())
