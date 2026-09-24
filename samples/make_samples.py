#!/usr/bin/env python3
"""make_samples.py — 生成 ppt-intake skill 的端到端测试样例。

base.pptx       基准：16:9、蓝色商务风、微软雅黑、8 页（含图片页/数据页）
material_a.pptx 材料：4:3、红色风、宋体+Arial、6 页（含表格页/图片页），
                  内容与基准部分重叠（数据更新）、部分新增、一处数据口径冲突
刻意让两者格式全面不一致，模拟"别人的 PPT"。
"""
import os
import struct
import zlib

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Cm, Pt

HERE = os.path.dirname(os.path.abspath(__file__))

YAHEI = "微软雅黑"
SONGTI = "宋体"


def make_png(path, w, h, top, bottom):
    """纯 stdlib 生成双色条带 PNG（模拟图表图片）。"""
    def chunk(tag, data):
        c = tag + data
        return struct.pack(">I", len(data)) + c + struct.pack(">I", zlib.crc32(c))

    rows = b""
    for y in range(h):
        color = top if (y // (h // 6)) % 2 == 0 else bottom
        px = b"\x00" + bytes(color) * w
        rows += px
    ihdr = struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0)
    with open(path, "wb") as f:
        f.write(b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", ihdr)
                + chunk(b"IDAT", zlib.compress(rows)) + chunk(b"IEND", b""))


def box(slide, x, y, w, h, text, size, color, bold=False, font=YAHEI,
        align=PP_ALIGN.LEFT, fill=None):
    from pptx.util import Emu
    tb = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    tf = tb.text_frame
    tf.word_wrap = True
    lines = text.split("\n")
    tf.text = lines[0]
    for line in lines[1:]:
        tf.add_paragraph().text = line  # 简单多行；样例无需保格式
    for p in tf.paragraphs:
        p.alignment = align
        for r in p.runs:
            r.font.name = font
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = color
    if fill:
        from pptx.enum.shapes import MSO_SHAPE
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(x - 0.5), Cm(y - 0.5),
                                    Cm(w + 1), Cm(h + 1))
        bg.fill.solid()
        bg.fill.fore_color.rgb = fill
        bg.line.fill.background()
        slide.shapes._spTree.remove(bg._element)
        slide.shapes._spTree.insert(2, bg._element)  # 垫底
    return tb


BLUE = RGBColor(0x17, 0x45, 0x7E)
NAVY = RGBColor(0x0E, 0x2A, 0x52)
DARK = RGBColor(0x33, 0x33, 0x33)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RED = RGBColor(0xC0, 0x00, 0x00)


def build_base():
    prs = Presentation()
    prs.slide_width = Cm(33.87)   # 16:9
    prs.slide_height = Cm(19.05)
    blank = prs.slide_layouts[6]
    s = []

    def add():
        s.append(prs.slides.add_slide(blank))
        return s[-1]

    # 0 cover
    sl = add()
    bg = sl.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid(); bg.fill.fore_color.rgb = NAVY; bg.line.fill.background()
    box(sl, 3, 7, 28, 3, "2025 年度经营汇报", 40, WHITE, bold=True)
    box(sl, 3, 11, 28, 2, "市场与竞争格局分析", 18, RGBColor(0xBF, 0xD3, 0xEE))
    # 1 toc
    sl = add()
    box(sl, 2.5, 1.5, 10, 2, "目录", 28, BLUE, bold=True)
    box(sl, 3.5, 5, 25, 9, "01 市场概览\n02 竞争格局\n03 总结与展望", 20, DARK)
    # 2 section 01
    sl = add()
    box(sl, 3, 6, 8, 4, "01", 54, BLUE, bold=True)
    box(sl, 8, 7.5, 15, 3, "市场概览", 32, DARK, bold=True)
    # 3 content 市场概览
    sl = add()
    box(sl, 2.5, 1.5, 20, 2, "市场概览", 24, BLUE, bold=True)
    box(sl, 3.5, 5, 26, 10, "市场规模 1.2 亿元\n同比增速 15%\n华东区域占比 40%", 18, DARK)
    # 4 stats
    sl = add()
    box(sl, 2.5, 1.5, 20, 2, "市场份额", 24, BLUE, bold=True)
    box(sl, 4, 5, 12, 5, "18%", 60, BLUE, bold=True)
    box(sl, 4, 11, 18, 2, "我们的市场份额（2025）", 14, DARK)
    box(sl, 4, 13.5, 18, 1.5, "数据来源：年度审计", 12, RGBColor(0x88, 0x88, 0x88))
    # 5 section 02
    sl = add()
    box(sl, 3, 6, 8, 4, "02", 54, BLUE, bold=True)
    box(sl, 8, 7.5, 15, 3, "竞争格局", 32, DARK, bold=True)
    # 6 content + picture 产品与竞品
    sl = add()
    box(sl, 2.5, 1.5, 20, 2, "产品与竞品", 24, BLUE, bold=True)
    box(sl, 3.5, 5, 14, 10, "旗舰产品 X1 系列\n竞品 A 售价 99 元\n渠道以直营为主", 18, DARK)
    png = os.path.join(HERE, "base_pic.png")
    make_png(png, 160, 120, (0x17, 0x45, 0x7E), (0xBF, 0xD3, 0xEE))
    sl.shapes.add_picture(png, Cm(19), Cm(5), Cm(12), Cm(9))
    # 7 closing
    sl = add()
    box(sl, 3, 6, 28, 4, "谢谢", 40, BLUE, bold=True, align=PP_ALIGN.CENTER)
    box(sl, 3, 12, 28, 2, "contact@example.com", 14, DARK, align=PP_ALIGN.CENTER)

    out = os.path.join(HERE, "base.pptx")
    prs.save(out)
    print("OK", out, len(prs.slides), "slides")


def build_material():
    prs = Presentation()  # 保持默认 4:3（10x7.5in）—— 刻意的尺寸冲突
    blank = prs.slide_layouts[6]

    def add():
        return prs.slides.add_slide(blank)

    # 0 cover（红）
    sl = add()
    bg = sl.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid(); bg.fill.fore_color.rgb = RED; bg.line.fill.background()
    box(sl, 1.5, 4, 22, 3, "XX 行业市场洞察 2026", 36, WHITE, bold=True, font=SONGTI)
    box(sl, 1.5, 8, 22, 2, "Third-Party Research", 16, WHITE, font="Arial")
    # 1 最新市场数据（与 base#3 重叠且更新）
    sl = add()
    box(sl, 1.5, 1, 15, 2, "最新市场数据", 24, RED, bold=True, font=SONGTI)
    box(sl, 2, 3.5, 20, 6, "市场规模 1.35 亿元（2025 实绩）\n同比增速 18%\n行业前二合计份额 45%",
        18, DARK, font=SONGTI)
    # 2 竞品对比表（新增；与 base#6 的"竞品A 99元"冲突 89 元）
    sl = add()
    box(sl, 1.5, 1, 15, 2, "主要竞品对比", 24, RED, bold=True, font=SONGTI)
    rows, cols = 4, 3
    tbl = sl.shapes.add_table(rows, cols, Cm(2), Cm(4), Cm(20), Cm(10)).table
    data = [["竞品", "售价", "份额"], ["竞品 A", "89 元", "12%"],
            ["竞品 B", "120 元", "9%"], ["竞品 C", "75 元", "6%"]]
    for r in range(rows):
        for c in range(cols):
            cell = tbl.cell(r, c)
            cell.text = data[r][c]
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    run.font.name = SONGTI
                    run.font.size = Pt(16)
                    run.font.color.rgb = WHITE if r == 0 else DARK
                    if r == 0:
                        run.font.bold = True
    # 3 渠道结构图（新增图片）
    sl = add()
    box(sl, 1.5, 1, 15, 2, "渠道结构", 24, RED, bold=True, font=SONGTI)
    png = os.path.join(HERE, "material_pic.png")
    make_png(png, 160, 120, (0xC0, 0x00, 0x00), (0xF2, 0xC8, 0xC8))
    sl.shapes.add_picture(png, Cm(6), Cm(4), Cm(14), Cm(10))
    # 4 政策风险（新增）
    sl = add()
    box(sl, 1.5, 1, 15, 2, "政策风险提示", 24, RED, bold=True, font=SONGTI)
    box(sl, 2, 3.5, 20, 6, "新版能效标准 2026Q2 实施\n补贴政策退坡 30%", 18, DARK, font=SONGTI)
    # 5 closing
    sl = add()
    box(sl, 1.5, 4, 22, 3, "THANKS", 44, RED, bold=True, font="Arial",
        align=PP_ALIGN.CENTER)

    out = os.path.join(HERE, "material_a.pptx")
    prs.save(out)
    print("OK", out, len(prs.slides), "slides")


if __name__ == "__main__":
    build_base()
    build_material()
